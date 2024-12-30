"""
Generate economic scenarios and export to Excel.
"""
import os
import sys
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from dataclasses import dataclass
from openpyxl.styles import Font, PatternFill, Alignment

# Add the project root directory to Python path
project_root = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, project_root)

from src.economic_scenario_generator import EconomicScenarioGenerator

@dataclass
class ScenarioConfig:
    """Configuration for economic scenario generation."""
    short_rate_mean: float
    short_rate_speed: float
    short_rate_vol: float
    equity_return_mean: float
    equity_vol: float
    jump_intensity: float
    jump_mean: float
    jump_vol: float
    credit_spread_mean: float
    credit_spread_vol: float
    inflation_mean: float
    inflation_vol: float

def generate_and_export_scenarios(num_scenarios=1000, projection_years=30):
    """Generate economic scenarios and export to Excel with detailed formatting."""
    
    # Initialize scenario generator with default parameters
    config = ScenarioConfig(
        short_rate_mean=0.02,  # 2% mean short rate
        short_rate_speed=0.2,  # Mean reversion speed
        short_rate_vol=0.01,   # Short rate volatility
        equity_return_mean=0.08,  # 8% mean equity return
        equity_vol=0.15,       # 15% equity volatility
        jump_intensity=0.1,    # Jump frequency
        jump_mean=-0.05,       # Average jump size
        jump_vol=0.02,         # Jump size volatility
        credit_spread_mean=0.01,  # 1% mean credit spread
        credit_spread_vol=0.005,  # Credit spread volatility
        inflation_mean=0.02,   # 2% mean inflation
        inflation_vol=0.01     # Inflation volatility
    )
    
    generator = EconomicScenarioGenerator(config)
    
    # Generate scenarios
    scenarios = generator.generate_scenarios(
        num_scenarios=num_scenarios,
        projection_years=projection_years
    )
    
    # Create timestamps for the projection period
    start_date = datetime.now()
    dates = [start_date + timedelta(days=i*365) for i in range(projection_years + 1)]
    
    # Create separate DataFrames for each economic factor
    dfs = {}
    factor_names = ['short_rate', 'long_rate', 'equity_return', 'credit_spread', 'inflation']
    
    for factor in factor_names:
        df = pd.DataFrame(
            scenarios[factor],
            index=dates,
            columns=[f'Scenario_{i+1}' for i in range(num_scenarios)]
        )
        dfs[factor] = df
    
    # Export to Excel with formatting
    output_file = 'economic_scenarios.xlsx'
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # Write summary statistics
        summary_stats = pd.DataFrame()
        for factor in factor_names:
            data = dfs[factor]
            stats = pd.DataFrame({
                'Mean': data.mean(axis=1).mean(),
                'Std Dev': data.std(axis=1).mean(),
                'Min': data.min().min(),
                'Max': data.max().max(),
                '5th Percentile': data.quantile(0.05).mean(),
                '95th Percentile': data.quantile(0.95).mean()
            }, index=[factor.replace('_', ' ').title()])
            summary_stats = pd.concat([summary_stats, stats])
        
        summary_stats.to_excel(writer, sheet_name='Summary Statistics')
        
        # Write detailed scenarios
        for factor in factor_names:
            sheet_name = factor.replace('_', ' ').title()
            dfs[factor].to_excel(writer, sheet_name=sheet_name)
            
            # Get the worksheet
            worksheet = writer.sheets[sheet_name]
            
            # Format headers
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
            
            for col in range(1, len(dfs[factor].columns) + 2):
                cell = worksheet.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
            
            # Format dates
            date_alignment = Alignment(horizontal='left')
            for row in range(2, len(dates) + 2):
                cell = worksheet.cell(row=row, column=1)
                cell.number_format = 'YYYY-MM-DD'
                cell.alignment = date_alignment
            
            # Format numbers
            number_alignment = Alignment(horizontal='right')
            for row in range(2, len(dates) + 2):
                for col in range(2, num_scenarios + 2):
                    cell = worksheet.cell(row=row, column=col)
                    cell.number_format = '0.00%'
                    cell.alignment = number_alignment
            
            # Adjust column widths
            for col in worksheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column].width = adjusted_width
    
    print(f"Economic scenarios have been exported to {output_file}")
    return output_file

if __name__ == '__main__':
    # Generate 1000 scenarios for 30 years
    output_file = generate_and_export_scenarios(num_scenarios=1000, projection_years=30)
